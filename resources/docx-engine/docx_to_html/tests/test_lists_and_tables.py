"""列表（有序/无序 + 层级）与表格（表头 + 边框 + 合并告警）。"""
from __future__ import annotations

import json

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_num_pr(paragraph, num_id: str, ilvl: str | None = None) -> None:
    """给段落挂上直接的编号属性（模拟 Word 界面创建的列表）。"""
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    if ilvl is not None:
        ilvl_el = OxmlElement("w:ilvl")
        ilvl_el.set(qn("w:val"), ilvl)
        num_pr.append(ilvl_el)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), num_id)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def test_style_based_lists_become_ul_and_ol(make_docx, extract_docx, tmp_path, read_html):
    """列表来自**样式**（python-docx / 部分导出器就这样）时也必须认得出来。"""

    def build(document):
        document.add_paragraph("无序项", style="List Bullet")
        document.add_paragraph("有序项一", style="List Number")
        document.add_paragraph("有序项二", style="List Number")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<ul><li>无序项</li></ul>" in html
    assert "<ol><li>有序项一</li><li>有序项二</li></ol>" in html


def test_direct_numbering_with_levels_nests_inside_previous_item(
    make_docx, extract_docx, tmp_path, read_html
):
    """直接编号属性 + ilvl → 嵌套列表；子列表必须落在上一个 <li> 里（HTML 合法嵌套）。

    numId=5 是 python-docx 默认模板里的十进制编号定义（见 requirements 里 python-docx<2 的约束）。
    """

    def build(document):
        parent = document.add_paragraph("一级项")
        _set_num_pr(parent, "5", "0")
        child = document.add_paragraph("二级项")
        _set_num_pr(child, "5", "1")
        sibling = document.add_paragraph("一级项二")
        _set_num_pr(sibling, "5", "0")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<ol><li>一级项<ol><li>二级项</li></ol></li><li>一级项二</li></ol>" in html


def test_unknown_numbering_definition_degrades_to_flat_ul_with_warning(
    make_docx, extract_docx, tmp_path, read_html
):
    def build(document):
        paragraph = document.add_paragraph("编号定义缺失的项")
        _set_num_pr(paragraph, "999", "3")

    docx_path = make_docx(build)
    result = extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    # 退化为一层无序列表（层级读不出来时不硬编多级）
    assert "<ul><li>编号定义缺失的项</li></ul>" in html
    assert any("编号定义缺失" in warning for warning in result["warnings"])


def test_table_first_row_becomes_th_and_border_from_cells(make_docx, extract_docx, tmp_path, read_html):
    def build(document):
        table = document.add_table(rows=2, cols=3)
        table.style = "Table Grid"
        for column, title in enumerate(("指标", "本周", "环比")):
            table.cell(0, column).text = title
        for column, value in enumerate(("转换成功率", "98.5%", "+1.2%")):
            table.cell(1, column).text = value

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<thead><tr><th" in html and "指标" in html and "环比" in html
    assert "<tbody>" in html and "<td" in html and "转换成功率" in html
    assert "border-collapse: collapse" in html
    assert "border: 1px solid #000000" in html  # Table Grid 的边框（w:color=auto → 黑）


def test_table_without_border_has_no_border_css(make_docx, extract_docx, tmp_path, read_html):
    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "甲"
        table.cell(1, 1).text = "乙"

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<table style=\"border-collapse: collapse;\">" in html
    assert "solid" not in html


def test_empty_cell_keeps_placeholder(make_docx, extract_docx, tmp_path, read_html):
    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "有内容"
        table.cell(1, 1).text = "也有内容"

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<th>&nbsp;</th>" in html or "<td>&nbsp;</td>" in html


def test_merged_cells_report_warning(make_docx, cli, tmp_path):
    def build(document):
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).merge(table.cell(0, 1))

    docx_path = make_docx(build)
    result = cli("extract", str(docx_path), "-o", str(tmp_path / "out.html"))
    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)

    assert any("合并单元格" in warning for warning in payload["warnings"])
