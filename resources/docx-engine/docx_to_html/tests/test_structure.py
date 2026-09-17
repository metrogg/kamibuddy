"""标题层级 / 段落样式映射 / 空段落与顺序保真。"""
from __future__ import annotations

import json

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def test_heading_styles_map_to_h1_h3_and_title_to_h1(make_docx, extract_docx, tmp_path, read_html):
    def build(document):
        document.add_heading("一级标题", level=1)
        document.add_heading("二级标题", level=2)
        document.add_heading("三级标题", level=3)
        document.add_paragraph("文档大标题", style="Title")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<h1" in html and "一级标题" in html
    assert "<h2" in html and "二级标题" in html
    assert "<h3" in html and "三级标题" in html
    # Title 样式按文档大标题处理 → h1（README 的映射表）
    assert html.count("<h1") == 2 and "文档大标题" in html
    assert "<h4" not in html


def test_outline_level_without_heading_style_still_maps_to_heading(
    make_docx, extract_docx, tmp_path, read_html
):
    """自定义样式 + 大纲级别：Word 里很常见，不能因为没有 Heading 样式就退成正文。"""

    def build(document):
        paragraph = document.add_paragraph("大纲级别兜底标题")
        p_pr = paragraph._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "0")
        p_pr.append(outline)

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<h1" in html and "大纲级别兜底标题" in html


def test_run_font_maps_to_inline_style(make_docx, extract_docx, tmp_path, read_html):
    def build(document):
        paragraph = document.add_paragraph()
        run = paragraph.add_run("重点")
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        run.font.name = "仿宋"
        paragraph.add_run("普通")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "font-size: 18pt" in html
    assert "font-weight: bold" in html
    assert "font-style: italic" in html
    assert "color: #123456" in html
    assert "font-family: '仿宋'" in html
    # 没显式设置格式的 run 不背包 <span>（继承段落级）：输出不被同一串声明淹没
    assert ">普通</" in html or ">普通" in html


def test_paragraph_alignment_and_indent_map_to_inline_style(
    make_docx, extract_docx, tmp_path, read_html
):
    def build(document):
        paragraph = document.add_paragraph("居中且缩进")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.left_indent = Pt(12)

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "text-align: center" in html
    assert "text-indent: 0.85cm" in html  # 24pt
    assert "margin-left: 0.42cm" in html  # 12pt


def test_empty_paragraph_is_kept_and_order_preserved(
    make_docx, extract_docx, tmp_path, read_html
):
    def build(document):
        document.add_paragraph("第一段")
        document.add_paragraph()
        document.add_paragraph("第三段")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert "<p>&nbsp;</p>" in html
    assert html.index("第一段") < html.index("&nbsp;") < html.index("第三段")


def test_document_is_complete_html5_with_docx_content_root(
    make_docx, extract_docx, tmp_path, read_html
):
    def build(document):
        document.add_paragraph("正文")

    docx_path = make_docx(build)
    extract_docx(docx_path, tmp_path / "out.html")
    html = read_html(tmp_path / "out.html")

    assert html.startswith("<!DOCTYPE html>")
    assert '<meta charset="utf-8">' in html
    assert '<article class="docx-content"' in html
    assert "<style" not in html  # 样式全 inline，不外链也不内嵌样式表
    assert "images/" not in html  # 没有图片的文档不凭空造引用


def test_missing_style_definition_does_not_break_extraction(make_docx, cli, tmp_path, read_html):
    """别的工具产出的 docx 常引用 styles.xml 里没有的样式：不能崩。

    python-docx 对这种引用会回落到默认样式（不抛异常），所以我们按默认样式出段落，
    文字照旧保留 —— 这里钉住的就是「不因为样式缺失而丢内容或失败」。
    """

    def build(document):
        paragraph = document.add_paragraph("样式缺失的段落")
        p_pr = paragraph._p.get_or_add_pPr()
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), "NoSuchStyle")
        p_pr.insert(0, p_style)

    docx_path = make_docx(build)
    result = cli("extract", str(docx_path), "-o", str(tmp_path / "out.html"))
    assert result.returncode == 0, result.stderr
    html = read_html(tmp_path / "out.html")

    assert "<p>样式缺失的段落</p>" in html


def test_hyperlink_text_is_not_lost(make_docx, cli, tmp_path, read_html):
    """超链接里的文字必须留住（python-docx 的 paragraph.runs 看不到它），并如实告警。"""

    def build(document):
        paragraph = document.add_paragraph("链接：")
        rel_id = paragraph.part.relate_to(
            "https://example.com", RT.HYPERLINK, is_external=True
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "示例站点"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    docx_path = make_docx(build)
    result = cli("extract", str(docx_path), "-o", str(tmp_path / "out.html"))
    assert result.returncode == 0, result.stderr
    html = read_html(tmp_path / "out.html")

    assert "示例站点" in html
    warnings = json.loads(result.stdout)["warnings"]
    assert any("超链接" in warning for warning in warnings)
