"""not_restorable 的**诚实性**：命中才报，不命中不报。

这是 spec 的硬要求：固定吐一串名字等于骗人，所以这里正反两个方向都测。
"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _extract(cli, docx_path: Path, tmp_path: Path) -> dict:
    result = cli("extract", str(docx_path), "-o", str(tmp_path / "out.html"))
    assert result.returncode == 0, result.stderr
    return json.loads(result.stdout)


def test_plain_document_reports_nothing(make_docx, cli, tmp_path):
    """干净文档必须报空 —— 这是诚实性的基准用例。"""

    def build(document):
        document.add_heading("标题", level=1)
        document.add_paragraph("正文")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "甲"

    docx_path = make_docx(build)
    assert _extract(cli, docx_path, tmp_path)["not_restorable"] == []


def test_empty_header_is_not_reported(make_docx, cli, tmp_path):
    """有 header 部件但没内容（只有空段落）→ 不算「有页眉」。"""

    def build(document):
        document.add_paragraph("正文")
        header = document.sections[0].header
        header.paragraphs[0].text = "先写上"
        header.paragraphs[0].text = ""  # 清空：部件留下了，内容没了

    docx_path = make_docx(build)
    assert _extract(cli, docx_path, tmp_path)["not_restorable"] == []


def test_header_with_text_reports_header_footer(make_docx, cli, tmp_path):
    def build(document):
        document.add_paragraph("正文")
        document.sections[0].header.paragraphs[0].text = "页眉文字"

    docx_path = make_docx(build)
    reported = _extract(cli, docx_path, tmp_path)["not_restorable"]
    assert "页眉页脚" in reported
    assert "页码" not in reported


def test_footer_with_page_field_reports_page_number(make_docx, cli, tmp_path):
    def build(document):
        document.add_paragraph("正文")
        footer = document.sections[0].footer
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " PAGE  \\* MERGEFORMAT ")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        footer.paragraphs[0]._p.append(field)

    docx_path = make_docx(build)
    assert "页码" in _extract(cli, docx_path, tmp_path)["not_restorable"]


def test_pageref_is_not_page_number(make_docx, cli, tmp_path):
    """PAGEREF（交叉引用）不是页码：只看域指令的第一个词。"""

    def build(document):
        document.add_paragraph("正文")
        footer = document.sections[0].footer
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " PAGEREF _Toc123 \\h ")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "3"
        run.append(text)
        field.append(run)
        footer.paragraphs[0]._p.append(field)

    docx_path = make_docx(build)
    reported = _extract(cli, docx_path, tmp_path)["not_restorable"]
    assert "页码" not in reported
    assert "页眉页脚" in reported  # 页脚本身有文字，仍然是丢了的东西


def test_section_break_and_body_field_are_reported(make_docx, cli, tmp_path):
    def build(document):
        document.add_paragraph("第一页")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("第二页")
        paragraph = document.add_paragraph()
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " DATE ")
        paragraph._p.append(field)

    docx_path = make_docx(build)
    reported = _extract(cli, docx_path, tmp_path)["not_restorable"]
    assert "分节" in reported
    assert "域代码" in reported


def test_floating_object_and_chart_are_reported(make_docx, cli, tmp_path):
    def build(document):
        paragraph = document.add_paragraph("带浮动对象")
        paragraph._p.append(OxmlElement("w:pict"))  # VML 文本框/形状的宿主元素

    docx_path = make_docx(build)
    # 图表：zip 里有 word/charts/ 才算（python-docx 造不出图表，直接补一个部件条目）
    with zipfile.ZipFile(docx_path, "a") as archive:
        archive.writestr("word/charts/chart1.xml", "<chart/>")

    reported = _extract(cli, docx_path, tmp_path)["not_restorable"]
    assert "浮动对象（文本框、形状）" in reported
    assert "图表" in reported


def test_footnote_and_comment_parts_are_reported_as_supplementary(make_docx, cli, tmp_path):
    def build(document):
        document.add_paragraph("正文")

    docx_path = make_docx(build)
    with zipfile.ZipFile(docx_path, "a") as archive:
        archive.writestr(
            "word/footnotes.xml",
            '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:footnote w:id="-1" w:type="separator"/><w:footnote w:id="1"><w:p><w:r><w:t>脚注</w:t>'
            "</w:r></w:p></w:footnote></w:footnotes>",
        )

    reported = _extract(cli, docx_path, tmp_path)["not_restorable"]
    assert "脚注/尾注" in reported
