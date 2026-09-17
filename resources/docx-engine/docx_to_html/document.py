"""document.py — 正文遍历（段落 / 表格按原顺序）与 HTML5 文档组装。

两个「为什么」：

1. **遍历用 ``document.iter_inner_content()``**：它按 body 里的真实顺序给出段落与
   表格；自己只扫 ``document.paragraphs`` 会把表格全甩到文末，段落顺序一乱版式就无从复用。
2. **列表要攒着出**：docx 的列表项是散落的普通段落，只有连着看才知道哪几段是同一个列表
   （见 lists.ListBuilder），所以这里遇到非列表块就把攒下的列表 flush 出去。

输出形状（下游对准这个契约）：完整 HTML5 文档 + 根内容元素
``<article class="docx-content">``（与 WorkBuddy format-extract 同约定，
便于 typeset / design-token 对接），样式全 inline。
"""
from __future__ import annotations

from html import escape
from pathlib import Path
from typing import TYPE_CHECKING

from docx.table import Table
from docx.text.paragraph import Paragraph

from .inline_style import (
    document_base_font,
    font_declarations,
    paragraph_declarations,
    style_attr,
)
from .lists import ListBuilder
from .paragraphs import (
    collect_runs,
    has_visual_content,
    render_block,
    render_inline_content,
)
from .tables import render_table

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

    from .context import RenderContext

DEFAULT_LANG = "zh-CN"
CONTENT_CLASS = "docx-content"


def build_html(document: "DocxDocument", ctx: "RenderContext", fallback_title: str) -> str:
    """整篇 docx → 一份完整 HTML5 文档字符串。"""
    blocks = _render_body(document, ctx)
    return _wrap(blocks, document, fallback_title)


def _render_body(document: "DocxDocument", ctx: "RenderContext") -> list[str]:
    parts: list[str] = []
    builder = ListBuilder()

    def flush() -> None:
        if builder:
            parts.append(builder.drain())

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            info = ctx.numbering.item_info(block, ctx)
            if info is not None:
                builder.add(info, *_list_item(block, ctx))
                continue
            flush()
            parts.append(render_block(block, ctx))
            continue
        if isinstance(block, Table):
            flush()
            parts.append(render_table(block, ctx))
    flush()
    return parts


def _list_item(paragraph: Paragraph, ctx: "RenderContext") -> tuple[str, str]:
    """列表项内容与它自己的段落样式（顺序与 ListBuilder.add 的参数一致）。"""
    runs = collect_runs(paragraph)
    content = (
        render_inline_content(paragraph, ctx, effective_fonts=False)
        if has_visual_content(paragraph, runs)
        else ""
    )
    style = style_attr(paragraph_declarations(paragraph, ctx))
    return content, style


def _wrap(blocks: list[str], document: "DocxDocument", fallback_title: str) -> str:
    # 文档基准字体（Normal 样式）挂在 <article> 上，作为整篇的兜底，
    # 这样段落级只需声明自己样式链改过的部分。
    base_style = style_attr(font_declarations(document_base_font(document)))
    article_attr = f'<article class="{CONTENT_CLASS}"'
    if base_style:
        article_attr += f' style="{base_style}"'
    article_attr += ">"
    body = "\n".join(blocks)
    return (
        "<!DOCTYPE html>\n"
        f'<html lang="{_language(document)}">\n'
        "<head>\n"
        '<meta charset="utf-8">\n'
        f"<title>{escape(_title(document, fallback_title))}</title>\n"
        "</head>\n"
        "<body>\n"
        f"{article_attr}\n"
        f"{body}\n"
        "</article>\n"
        "</body>\n"
        "</html>\n"
    )


def _title(document: "DocxDocument", fallback_title: str) -> str:
    core_title = document.core_properties.title
    if core_title is not None and core_title.strip():
        return core_title.strip()
    return fallback_title


def _language(document: "DocxDocument") -> str:
    language = document.core_properties.language
    if language is not None and language.strip():
        return language.strip()
    return DEFAULT_LANG


def default_title_for(docx_path: str) -> str:
    """标题兜底：用文件名（不含扩展名）。"""
    return Path(docx_path).stem
