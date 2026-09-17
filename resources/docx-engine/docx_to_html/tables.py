"""tables.py — 表格 → HTML（``thead``/``th`` + 单元格文本 + 基础边框）。

为什么表头按「首行」判：docx 里没有「表头」这个概念（``w:tblHeader`` 只表示
跨页重复，并不保证是第一行），而 HTML 侧下游（typeset / 预览）需要 ``th``；
按首行处理是与 WorkBuddy format-extract 同款的约定，也是报表类文档的实际情形。

为什么边框要三处找：Word 把表格边框分散在表级 ``w:tblBorders``、表样式定义、
单元格 ``w:tcBorders`` 三处（正向引擎就走第三条），按此优先级取第一处有值的，
并转成**单元格上**的 border 声明 —— HTML 的表格边框必须画在单元格上才看得见。
"""
from __future__ import annotations

from typing import TYPE_CHECKING

from docx.oxml.ns import qn

from .inline_style import paragraph_declarations, style_attr
from .paragraphs import render_inline_content
from .style_map import border_width_px

if TYPE_CHECKING:
    from docx.table import Table, _Cell

    from .context import RenderContext

# 取边框时按这个顺序看（先外框后内框），取到第一个有效值即用。
_BORDER_TAGS = ("top", "left", "bottom", "right", "insideH", "insideV")
_NO_BORDER_VALUES = frozenset({"none", "nil"})


def render_table(table: "Table", ctx: "RenderContext") -> str:
    """整张表 → HTML；首行进 ``thead``（``th``），其余进 ``tbody``（``td``）。"""
    _warn_on_merged_cells(table, ctx)
    if any(cell.tables for row in table.rows for cell in row.cells):
        ctx.warn("表格单元格内的嵌套表格未提取（只保留单元格文本）")
    border_css = _border_css(table)
    parts: list[str] = ['<table style="border-collapse: collapse;">']
    rows = list(table.rows)
    if rows:
        parts.append("<thead><tr>")
        for cell in rows[0].cells:
            parts.append(_cell_html(cell, ctx, "th", border_css))
        parts.append("</tr></thead>")
    if len(rows) > 1:
        parts.append("<tbody>")
        for row in rows[1:]:
            parts.append("<tr>")
            for cell in row.cells:
                parts.append(_cell_html(cell, ctx, "td", border_css))
            parts.append("</tr>")
        parts.append("</tbody>")
    parts.append("</table>")
    return "".join(parts)


def _cell_html(cell: "_Cell", ctx: "RenderContext", tag: str, border_css: str | None) -> str:
    paragraphs = cell.paragraphs
    declarations = dict(paragraph_declarations(paragraphs[0], ctx)) if paragraphs else {}
    if border_css:
        declarations["border"] = border_css
    style = style_attr(declarations)
    style_attr_html = f' style="{style}"' if style else ""
    inner = "<br>".join(
        render_inline_content(paragraph, ctx, effective_fonts=True) for paragraph in paragraphs
    )
    if inner.strip() == "":
        inner = "&nbsp;"
    return f"<{tag}{style_attr_html}>{inner}</{tag}>"


def _warn_on_merged_cells(table: "Table", ctx: "RenderContext") -> None:
    """合并单元格（gridSpan / vMerge）按展开后的矩形输出，文本会重复 —— 如实告警。

    直接在 XML 上看 gridSpan/vMerge，不能拿 ``row.cells`` 判重：python-docx 每次
    访问都会新建 lxml 代理对象，``id()`` 不稳定（同一单元格会判成不同的两个），
    会误报「有合并单元格」。
    """
    for tc in table._tbl.iter(qn("w:tc")):
        tc_pr = tc.find(qn("w:tcPr"))
        if tc_pr is None:
            continue
        if _grid_span(tc_pr) > 1 or tc_pr.find(qn("w:vMerge")) is not None:
            ctx.warn("表格存在合并单元格：未还原合并（按展开的矩形输出，文本可能重复）")
            return


def _grid_span(tc_pr) -> int:
    grid_span = tc_pr.find(qn("w:gridSpan"))
    if grid_span is None:
        return 1
    return _int_or_none(grid_span.get(qn("w:val"))) or 1


def _border_css(table: "Table") -> str | None:
    for borders in (
        _table_borders(table),
        _style_borders(table),
        _cell_borders(table),
    ):
        if borders is None:
            continue
        css = _borders_to_css(borders)
        if css is not None:
            return css
    return None


def _table_borders(table: "Table"):
    tbl_pr = table._tbl.find(qn("w:tblPr"))
    return None if tbl_pr is None else tbl_pr.find(qn("w:tblBorders"))


def _style_borders(table: "Table"):
    style = table.style
    if style is None:
        return None
    style_tbl_pr = style.element.find(qn("w:tblPr"))
    return None if style_tbl_pr is None else style_tbl_pr.find(qn("w:tblBorders"))


def _cell_borders(table: "Table"):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.find(qn("w:tcPr"))
            if tc_pr is None:
                continue
            borders = tc_pr.find(qn("w:tcBorders"))
            if borders is not None:
                return borders
    return None


def _borders_to_css(borders) -> str | None:
    for name in _BORDER_TAGS:
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            continue
        value = edge.get(qn("w:val"))
        if value is None or value in _NO_BORDER_VALUES:
            continue
        color = edge.get(qn("w:color"))
        hex_color = f"#{color}" if color is not None and color.lower() != "auto" else "#000000"
        return f"{border_width_px(_int_or_none(edge.get(qn('w:sz'))))}px solid {hex_color}"
    return None


def _int_or_none(raw: str | None) -> int | None:
    if raw is None:
        return None
    try:
        return int(raw)
    except ValueError:
        return None
